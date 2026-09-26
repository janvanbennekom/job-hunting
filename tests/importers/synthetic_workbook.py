"""Build small synthetic project spreadsheets for tests."""

from __future__ import annotations

from pathlib import Path

import openpyxl

# Avoid collision with real workbook rows / imported assignment IDs in shared PostgreSQL.
DATA_START_ROW = 500

SYNTH_SEQUENCE_ALPHA = 9001
SYNTH_SEQUENCE_BETA = 9002
SYNTH_SEQUENCE_DUP_ID = 9003

SYNTH_CODE_LA = "syn-fd-LA"
SYNTH_CODE_LADM = "syn-fda-LADM"
SYNTH_CODE_API = "syn-fsd-API"


def write_synthetic_workbook(
    path: Path,
    *,
    include_unknown_group: bool = False,
    duplicate_project_id: bool = False,
    unexpected_capability_value: bool = False,
) -> None:
    wb = openpyxl.Workbook()
    et = wb.active
    et.title = "ExperienceTypes"
    et.append(["filter_code", "group", "topic", "description", "Example"])
    et.append(
        [
            SYNTH_CODE_LA,
            "Domain",
            "Land Administration",
            "Land administration domain",
            "example",
        ]
    )
    et.append(
        [
            SYNTH_CODE_LADM,
            "Data and Analysis",
            "LADM",
            "ISO19152",
            None,
        ]
    )
    et.append(
        [
            SYNTH_CODE_API,
            "System Development",
            "System integration and interoperability",
            "Integration",
            None,
        ]
    )
    if include_unknown_group:
        et.append(["bad-1", "Unknown Group", "Bad capability", None, None])

    project = wb.create_sheet("Project")
    project.append(
        [
            "id",
            "sequence",
            "name",
            "assignment_name",
            "beneficiary",
            "donor",
            "country",
            "period",
            "working_days",
            "last_year",
            "project_description",
            "tools",
            "responsibilities",
            "web_link",
            "Land Administration",
            "LADM",
            "System integration and interoperability",
        ]
    )
    row = DATA_START_ROW
    project.cell(row=row, column=1, value="P-001")
    project.cell(row=row, column=2, value=SYNTH_SEQUENCE_ALPHA)
    project.cell(row=row, column=3, value="Project Alpha")
    project.cell(row=row, column=4, value="GIS lead")
    project.cell(row=row, column=5, value="Client A")
    project.cell(row=row, column=6, value="Donor A")
    project.cell(row=row, column=7, value="Country A")
    project.cell(row=row, column=8, value="2020-2021")
    project.cell(row=row, column=9, value=50)
    project.cell(row=row, column=10, value=2021)
    project.cell(row=row, column=11, value="Description A")
    project.cell(row=row, column=12, value="QGIS")
    project.cell(row=row, column=13, value="Lead GIS")
    project.cell(row=row, column=14, value="https://example.org/a")
    project.cell(row=row, column=15, value=SYNTH_CODE_LA)
    project.cell(row=row, column=16, value=SYNTH_CODE_LADM)
    project.cell(row=row, column=17, value=SYNTH_CODE_API)

    row2 = DATA_START_ROW + 1
    project.cell(row=row2, column=1, value="P-002")
    project.cell(row=row2, column=2, value=SYNTH_SEQUENCE_BETA)
    project.cell(row=row2, column=3, value="Project Beta")
    project.cell(row=row2, column=4, value="Advisor")
    project.cell(row=row2, column=5, value="Client B")
    project.cell(row=row2, column=7, value="Country B")
    project.cell(row=row2, column=8, value="2018-2019")
    project.cell(row=row2, column=9, value=30)
    project.cell(row=row2, column=10, value=2019)
    project.cell(row=row2, column=11, value="Description B")
    project.cell(row=row2, column=15, value=SYNTH_CODE_LA)

    if duplicate_project_id:
        row3 = DATA_START_ROW + 2
        project.cell(row=row3, column=1, value="P-001")
        project.cell(row=row3, column=2, value=SYNTH_SEQUENCE_DUP_ID)
        project.cell(row=row3, column=3, value="Duplicate ID project")
        project.cell(row=row3, column=4, value="Role")
        project.cell(row=row3, column=8, value="2022")
        project.cell(row=row3, column=9, value=10)
        project.cell(row=row3, column=10, value=2022)
        project.cell(row=row3, column=11, value="Dup")

    if unexpected_capability_value:
        project.cell(row=DATA_START_ROW, column=15, value="unexpected-value")

    wb.save(path)
    wb.close()
