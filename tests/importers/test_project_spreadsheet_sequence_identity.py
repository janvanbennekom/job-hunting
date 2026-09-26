"""Assignment identity and sequence validation tests."""

from pathlib import Path

import openpyxl
import pytest

from jobhunter.infrastructure.importers.project_spreadsheet.identity import (
    assignment_id,
)
from jobhunter.infrastructure.importers.project_spreadsheet.reader import read_workbook
from jobhunter.infrastructure.importers.project_spreadsheet.validation import (
    validate_snapshot,
)
from importers.synthetic_workbook import (
    SYNTH_SEQUENCE_ALPHA,
    SYNTH_SEQUENCE_BETA,
    write_synthetic_workbook,
)


def test_assignment_id_uses_source_reference_and_sequence(tmp_path: Path) -> None:
    ref = (tmp_path / "book.xlsx").as_posix()
    assert assignment_id(ref, 9001) == assignment_id(ref, 9001)
    assert assignment_id(ref, 9001) != assignment_id(ref, 9002)
    assert assignment_id(ref, 9001) != assignment_id("other/book.xlsx", 9001)


def test_reordering_rows_preserves_assignment_ids(tmp_path: Path) -> None:
    path = tmp_path / "reorder.xlsx"
    write_synthetic_workbook(path)
    ref = path.as_posix()
    id_alpha = assignment_id(ref, SYNTH_SEQUENCE_ALPHA)
    id_beta = assignment_id(ref, SYNTH_SEQUENCE_BETA)

    wb = openpyxl.load_workbook(path)
    project = wb["Project"]
    row_a = [project.cell(row=500, column=c).value for c in range(1, 18)]
    row_b = [project.cell(row=501, column=c).value for c in range(1, 18)]
    for c in range(1, 18):
        project.cell(row=500, column=c, value=row_b[c - 1])
        project.cell(row=501, column=c, value=row_a[c - 1])
    wb.save(path)
    wb.close()

    snapshot = read_workbook(path)
    assert assignment_id(ref, SYNTH_SEQUENCE_ALPHA) == id_alpha
    assert assignment_id(ref, SYNTH_SEQUENCE_BETA) == id_beta


def test_inserted_row_does_not_change_existing_assignment_ids(tmp_path: Path) -> None:
    path = tmp_path / "insert.xlsx"
    write_synthetic_workbook(path)
    ref = path.as_posix()
    before = {
        SYNTH_SEQUENCE_ALPHA: assignment_id(ref, SYNTH_SEQUENCE_ALPHA),
        SYNTH_SEQUENCE_BETA: assignment_id(ref, SYNTH_SEQUENCE_BETA),
    }

    wb = openpyxl.load_workbook(path)
    project = wb["Project"]
    project.insert_rows(2)
    project.cell(row=2, column=1, value="P-INSERT")
    project.cell(row=2, column=2, value=8999)
    project.cell(row=2, column=3, value="Inserted row")
    wb.save(path)
    wb.close()

    snapshot = read_workbook(path)
    assert len(snapshot.project_rows) == 3
    assert assignment_id(ref, SYNTH_SEQUENCE_ALPHA) == before[SYNTH_SEQUENCE_ALPHA]
    assert assignment_id(ref, SYNTH_SEQUENCE_BETA) == before[SYNTH_SEQUENCE_BETA]


def test_missing_sequence_is_validation_error(tmp_path: Path) -> None:
    path = tmp_path / "missing-seq.xlsx"
    write_synthetic_workbook(path)
    wb = openpyxl.load_workbook(path)
    wb["Project"].cell(row=500, column=2, value="")
    wb.save(path)
    wb.close()

    snapshot = read_workbook(path)
    errors, _warnings = validate_snapshot(snapshot, path)
    assert any("missing sequence" in e for e in errors)


def test_duplicate_sequence_is_validation_error(tmp_path: Path) -> None:
    path = tmp_path / "dup-seq.xlsx"
    write_synthetic_workbook(path)
    wb = openpyxl.load_workbook(path)
    wb["Project"].cell(row=501, column=2, value=SYNTH_SEQUENCE_ALPHA)
    wb.save(path)
    wb.close()

    snapshot = read_workbook(path)
    errors, _warnings = validate_snapshot(snapshot, path)
    assert any("Duplicate Project.sequence" in e for e in errors)


def test_duplicate_project_id_remains_warning_only(tmp_path: Path) -> None:
    path = tmp_path / "dup-id.xlsx"
    write_synthetic_workbook(path, duplicate_project_id=True)
    snapshot = read_workbook(path)
    errors, warnings = validate_snapshot(snapshot, path)
    assert not errors
    assert any("Duplicate Project.id" in w for w in warnings)
