"""Unit tests for project spreadsheet reading and validation."""

from pathlib import Path

import openpyxl
import pytest

from jobhunter.infrastructure.importers.project_spreadsheet.reader import read_workbook
from jobhunter.infrastructure.importers.project_spreadsheet.validation import (
    validate_snapshot,
)
from importers.synthetic_workbook import (
    SYNTH_CODE_API,
    SYNTH_CODE_LA,
    SYNTH_CODE_LADM,
    write_synthetic_workbook,
)


def test_read_synthetic_workbook_maps_capabilities(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.xlsx"
    write_synthetic_workbook(path)
    snapshot = read_workbook(path)
    assert len(snapshot.capability_definitions) == 3
    assert len(snapshot.project_rows) == 2
    alpha = snapshot.project_rows[0]
    assert alpha.source_record_id == "P-001"
    assert alpha.project_name == "Project Alpha"
    assert alpha.period_text == "2020-2021"
    assert alpha.working_days == 50
    assert SYNTH_CODE_LA in alpha.capability_codes
    assert SYNTH_CODE_LADM in alpha.capability_codes
    assert SYNTH_CODE_API in alpha.capability_codes
    beta = snapshot.project_rows[1]
    assert SYNTH_CODE_LADM not in beta.capability_codes


def test_new_experience_type_and_project_column_import_without_code_change(
    tmp_path: Path,
) -> None:
    path = tmp_path / "new-cap.xlsx"
    wb = openpyxl.Workbook()
    et = wb.active
    et.title = "ExperienceTypes"
    et.append(["filter_code", "group", "topic", "description", "Example"])
    et.append(
        [
            "syn-new-cap",
            "Domain",
            "Newly Added Capability",
            "Added only in workbook",
            None,
        ]
    )
    project = wb.create_sheet("Project")
    project.append(["id", "name", "Newly Added Capability"])
    project.append(["P-NEW", "Project with new cap", "syn-new-cap"])
    wb.save(path)
    wb.close()

    snapshot = read_workbook(path)
    assert len(snapshot.capability_definitions) == 1
    assert snapshot.capability_definitions[0].code == "syn-new-cap"
    assert len(snapshot.project_rows) == 1
    assert "syn-new-cap" in snapshot.project_rows[0].capability_codes


def test_abbreviated_project_header_matches_filter_code_suffix(tmp_path: Path) -> None:
    """Long topic in ExperienceTypes, short Project column from filter_code suffix."""
    path = tmp_path / "abbrev.xlsx"
    wb = openpyxl.Workbook()
    et = wb.active
    et.title = "ExperienceTypes"
    et.append(["filter_code", "group", "topic", "description", "Example"])
    et.append(
        [
            "fda-LADM",
            "Data and Analysis",
            "Land Administration Domain Model",
            "ISO19152",
            None,
        ]
    )
    project = wb.create_sheet("Project")
    project.append(["id", "name", "LADM"])
    project.append(["P-1", "Example project", "fda-LADM"])
    wb.save(path)
    wb.close()

    snapshot = read_workbook(path)
    assert len(snapshot.project_rows) == 1
    assert "fda-LADM" in snapshot.project_rows[0].capability_codes


def test_duplicate_project_id_is_validation_warning(tmp_path: Path) -> None:
    path = tmp_path / "dup.xlsx"
    write_synthetic_workbook(path, duplicate_project_id=True)
    snapshot = read_workbook(path)
    errors, warnings = validate_snapshot(snapshot, path)
    assert not errors
    assert any("Duplicate Project.id" in w for w in warnings)


def test_unknown_capability_group_is_validation_error(tmp_path: Path) -> None:
    path = tmp_path / "bad-group.xlsx"
    write_synthetic_workbook(path, include_unknown_group=True)
    snapshot = read_workbook(path)
    errors, _warnings = validate_snapshot(snapshot, path)
    assert any("Unknown capability group" in e for e in errors)


def test_unexpected_capability_cell_value_warns(tmp_path: Path) -> None:
    path = tmp_path / "bad-cell.xlsx"
    write_synthetic_workbook(path, unexpected_capability_value=True)
    snapshot = read_workbook(path)
    errors, warnings = validate_snapshot(snapshot, path)
    assert not errors
    assert any("unexpected capability cell value" in w for w in warnings)
