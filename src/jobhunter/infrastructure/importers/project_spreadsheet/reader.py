"""Read Project and ExperienceTypes worksheets from an XLSX workbook."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from jobhunter.infrastructure.importers.project_spreadsheet.types import (
    CapabilityColumn,
    CapabilityDefinitionRow,
    ProjectRow,
    WorkbookSnapshot,
)

PROJECT_SHEET = "Project"
EXPERIENCE_TYPES_SHEET = "ExperienceTypes"

PROJECT_FIELD_COLUMNS: dict[str, str] = {
    "id": "source_record_id",
    "sequence": "sequence",
    "name": "project_name",
    "assignment_name": "role",
    "beneficiary": "beneficiary",
    "donor": "donor",
    "country": "country",
    "period": "period_text",
    "working_days": "working_days",
    "last_year": "last_active_year",
    "project_description": "description",
    "tools": "tools_technologies_text",
    "responsibilities": "responsibilities",
    "web_link": "web_link",
}


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell_int(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    return int(str(value).strip())


def _interpret_capability_cell(value: Any, expected_code: str) -> tuple[bool, str | None]:
    """Return (active, warning). Active only when value unambiguously marks capability."""
    if value is None:
        return False, None
    if isinstance(value, bool):
        return value, None if value else None
    if isinstance(value, (int, float)):
        if value == 0:
            return False, None
        if value == 1:
            return True, None
        return False, f"unexpected numeric capability value {value!r}"

    text = str(value).strip()
    if not text:
        return False, None
    lower = text.lower()
    if lower in {"0", "false", "no"}:
        return False, None
    if lower == expected_code.lower():
        return True, None
    if lower in {"x", "1", "true", "yes"}:
        return True, None
    return False, f"unexpected capability cell value {text!r} (expected {expected_code})"


def build_header_to_capability_code(
    definitions: list[CapabilityDefinitionRow],
) -> dict[str, str]:
    """Map Project column headers to capability codes from ExperienceTypes.

    Primary match: column header equals capability topic.
    Secondary match: header equals the segment after the last hyphen in
    filter_code (abbreviated column titles in Project).
    """
    mapping: dict[str, str] = {}
    for definition in definitions:
        if definition.topic:
            mapping[definition.topic] = definition.code
    for definition in definitions:
        if "-" not in definition.code:
            continue
        alias = definition.code.rsplit("-", 1)[-1]
        if not alias:
            continue
        existing = mapping.get(alias)
        if existing is None:
            mapping[alias] = definition.code
        elif existing != definition.code:
            continue
    return mapping


def read_workbook(path: Path) -> WorkbookSnapshot:
    source_path = str(path.resolve())
    source_reference = path.as_posix()

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if EXPERIENCE_TYPES_SHEET not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet {EXPERIENCE_TYPES_SHEET!r}")
        if PROJECT_SHEET not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet {PROJECT_SHEET!r}")

        et_rows = list(workbook[EXPERIENCE_TYPES_SHEET].iter_rows(values_only=True))
        et_header = [_normalize_header(c) for c in et_rows[0]]
        code_idx = et_header.index("filter_code")
        group_idx = et_header.index("group")
        topic_idx = et_header.index("topic")
        desc_idx = et_header.index("description") if "description" in et_header else -1
        example_idx = et_header.index("Example") if "Example" in et_header else -1

        capability_definitions: list[CapabilityDefinitionRow] = []
        topic_to_code: dict[str, str] = {}
        for row in et_rows[1:]:
            if not row or row[code_idx] is None:
                continue
            code = _normalize_header(row[code_idx])
            if not code:
                continue
            group = _normalize_header(row[group_idx])
            topic = _normalize_header(row[topic_idx])
            description = (
                _cell_str(row[desc_idx]) if desc_idx >= 0 and desc_idx < len(row) else None
            )
            example = (
                _cell_str(row[example_idx])
                if example_idx >= 0 and example_idx < len(row)
                else None
            )
            capability_definitions.append(
                CapabilityDefinitionRow(
                    code=code,
                    group=group,
                    topic=topic,
                    description=description,
                    example=example,
                )
            )
        header_to_code = build_header_to_capability_code(capability_definitions)

        project_rows_raw = list(workbook[PROJECT_SHEET].iter_rows(values_only=True))
        project_header = [_normalize_header(c) for c in project_rows_raw[0]]
        header_index = {name: idx for idx, name in enumerate(project_header) if name}

        capability_columns: list[CapabilityColumn] = []
        for idx, header in enumerate(project_header):
            if not header:
                continue
            if header in header_to_code:
                capability_columns.append(
                    CapabilityColumn(
                        column_index=idx,
                        header=header,
                        capability_code=header_to_code[header],
                    )
                )

        project_rows: list[ProjectRow] = []
        for excel_row_number, row in enumerate(project_rows_raw[1:], start=2):
            if not row:
                continue
            id_idx = header_index.get("id")
            if id_idx is None or id_idx >= len(row) or row[id_idx] is None:
                continue
            source_record_id = _normalize_header(row[id_idx])
            if not source_record_id:
                continue

            name_idx = header_index["name"]
            project_name = _normalize_header(row[name_idx]) if name_idx < len(row) else ""
            if not project_name:
                project_name = ""

            active_codes: set[str] = set()
            for col in capability_columns:
                cell = row[col.column_index] if col.column_index < len(row) else None
                active, _warning = _interpret_capability_cell(cell, col.capability_code)
                if active:
                    active_codes.add(col.capability_code)

            project_rows.append(
                ProjectRow(
                    excel_row_number=excel_row_number,
                    source_record_id=source_record_id,
                    sequence=_cell_int(row[header_index["sequence"]])
                    if "sequence" in header_index
                    else None,
                    project_name=project_name,
                    role=_cell_str(row[header_index["assignment_name"]])
                    if "assignment_name" in header_index
                    else None,
                    beneficiary=_cell_str(row[header_index["beneficiary"]])
                    if "beneficiary" in header_index
                    else None,
                    donor=_cell_str(row[header_index["donor"]])
                    if "donor" in header_index
                    else None,
                    country=_cell_str(row[header_index["country"]])
                    if "country" in header_index
                    else None,
                    period_text=_cell_str(row[header_index["period"]])
                    if "period" in header_index
                    else None,
                    working_days=_cell_int(row[header_index["working_days"]])
                    if "working_days" in header_index
                    else None,
                    last_active_year=_cell_int(row[header_index["last_year"]])
                    if "last_year" in header_index
                    else None,
                    description=_cell_str(row[header_index["project_description"]])
                    if "project_description" in header_index
                    else None,
                    responsibilities=_cell_str(row[header_index["responsibilities"]])
                    if "responsibilities" in header_index
                    else None,
                    tools_technologies_text=_cell_str(row[header_index["tools"]])
                    if "tools" in header_index
                    else None,
                    web_link=_cell_str(row[header_index["web_link"]])
                    if "web_link" in header_index
                    else None,
                    capability_codes=frozenset(active_codes),
                )
            )
    finally:
        workbook.close()

    return WorkbookSnapshot(
        source_path=source_path,
        source_reference=source_reference,
        capability_definitions=capability_definitions,
        capability_columns=capability_columns,
        project_rows=project_rows,
    )
